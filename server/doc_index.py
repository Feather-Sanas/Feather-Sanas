"""
Document RAG over user-uploaded unstructured files (PDF / DOCX / TXT / MD).

Sani grounds chat answers in two corpora: the crawled sanas.ai site (webindex.py)
and the documents a user uploads here. This module parses an upload into text,
splits it into ~2 kB chunks, and indexes them with the SAME lexical term-frequency
scoring as webindex — so a query can retrieve from the user's own material
alongside the site, with no embeddings service and no per-request cost.

The store is PERSISTED TO DISK (server/rag_store.json) and shared across sessions,
so uploaded docs survive a backend restart. Parsing degrades gracefully: a scanned
(image-only) PDF yields no text and is reported back to the caller, not crashed on.
"""
from __future__ import annotations

import io
import json
import os
import re
import time
import uuid
from pathlib import Path

# Reuse the site retriever's stop-word list + tokenizer so doc scores are
# directly comparable to web scores when the two result sets are merged.
from webindex import _terms

# Store location: defaults to this module's dir for local dev; SAN_DATA_DIR lets a
# container point it at a mounted volume so uploads survive image rebuilds.
_DIR = Path(os.getenv("SAN_DATA_DIR", str(Path(__file__).resolve().parent)))
try:
    _DIR.mkdir(parents=True, exist_ok=True)
except Exception:
    _DIR = Path(__file__).resolve().parent
_PATH = _DIR / "rag_store.json"
CHUNK_CHARS = 2000          # ~500 tokens; long enough for lexical term overlap
MIN_SCORE = 2               # drop near-zero matches so a doc isn't forced in

# In-memory mirror of the on-disk store. Each chunk carries a precomputed
# lowercase copy ("_x") for scoring; that field is NOT persisted.
_docs: list[dict] = []
_loaded = False


# ---- persistence ------------------------------------------------------------
def _load() -> list[dict]:
    global _docs, _loaded
    if _loaded:
        return _docs
    _loaded = True
    try:
        data = json.loads(_PATH.read_text())
        _docs = data.get("docs", [])
        for d in _docs:
            for c in d.get("chunks", []):
                c["_x"] = c["text"].lower()
    except Exception:
        _docs = []
    return _docs


def _save() -> None:
    slim = {"docs": [
        {k: v for k, v in d.items() if k != "chunks"} |
        {"chunks": [{kk: vv for kk, vv in c.items() if kk != "_x"} for c in d["chunks"]]}
        for d in _docs
    ]}
    tmp = _PATH.with_suffix(".tmp")
    tmp.write_text(json.dumps(slim, ensure_ascii=False))
    tmp.replace(_PATH)


# ---- parsing ----------------------------------------------------------------
def _strip_html(s: str) -> str:
    s = re.sub(r"(?is)<(script|style)\b.*?</\1>", " ", s)
    s = re.sub(r"(?s)<[^>]+>", " ", s)
    s = (s.replace("&nbsp;", " ").replace("&amp;", "&").replace("&lt;", "<")
          .replace("&gt;", ">").replace("&#39;", "'").replace("&quot;", '"'))
    return s


def extract_text(filename: str, raw: bytes) -> str:
    """Best-effort plain-text extraction by extension / sniffed type."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or raw[:5] == b"%PDF-":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages)
    if name.endswith(".docx") or raw[:2] == b"PK":   # docx is a zip container
        from docx import Document
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs]
        for tbl in doc.tables:                        # flatten table cells to lines
            for row in tbl.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    text = raw.decode("utf-8", errors="replace")
    if name.endswith((".html", ".htm")) or "<html" in text[:2000].lower():
        text = _strip_html(text)
    return text


def _chunk(text: str) -> list[str]:
    """Split on blank lines, then pack paragraphs into <= CHUNK_CHARS blocks;
    over-long paragraphs are sentence-split so one wall of text still chunks."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: list[str] = []
    for p in paras:
        if len(p) <= CHUNK_CHARS:
            pieces.append(p)
        else:
            sent, buf = re.split(r"(?<=[.!?])\s+", p), ""
            for s in sent:
                if len(buf) + len(s) + 1 > CHUNK_CHARS and buf:
                    pieces.append(buf); buf = s
                else:
                    buf = (buf + " " + s).strip()
            if buf:
                pieces.append(buf)
    chunks, buf = [], ""
    for p in pieces:                                  # coalesce small paras together
        if len(buf) + len(p) + 2 > CHUNK_CHARS and buf:
            chunks.append(buf); buf = p
        else:
            buf = (buf + "\n\n" + p).strip()
    if buf:
        chunks.append(buf)
    return chunks


# ---- public API -------------------------------------------------------------
def ingest(filename: str, raw: bytes) -> dict:
    """Parse + chunk + index one upload; persist; return a summary dict.
    Raises ValueError if no extractable text (e.g. a scanned PDF)."""
    _load()
    text = extract_text(filename, raw).strip()
    if len(text) < 20:
        raise ValueError("no extractable text (the file may be scanned/image-only or empty)")
    doc_id = uuid.uuid4().hex[:12]
    chunks = []
    for i, body in enumerate(_chunk(text)):
        chunks.append({"chunk_id": i, "text": body, "_x": body.lower()})
    doc = {"doc_id": doc_id, "name": filename or f"document-{doc_id}",
           "chars": len(text), "added": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
           "chunks": chunks}
    _docs.append(doc)
    _save()
    return {"doc_id": doc_id, "name": doc["name"], "chunks": len(chunks),
            "chars": len(text), "total_docs": len(_docs)}


def search(query: str, k: int = 2, min_score: int = MIN_SCORE) -> list[dict]:
    """Lexical top-k over all uploaded chunks. Shape matches webindex.search()
    (title/url/snippet) plus kind='doc' and doc_name so callers can mark the
    source as user-supplied rather than a sanas.ai page."""
    docs, terms = _load(), _terms(query or "")
    if not docs or not terms:
        return []
    scored = []
    for d in docs:
        for c in d["chunks"]:
            s = sum(c["_x"].count(t) for t in terms)
            if s >= min_score:
                scored.append((s, d, c))
    scored.sort(key=lambda x: -x[0])
    out = []
    for _, d, c in scored[:k]:
        # Send the FULL matched chunk to the LLM as grounding context (chunks are
        # already capped at CHUNK_CHARS) — a short teaser would hide the answer.
        # The UI only renders the document name, never this text.
        out.append({"title": d["name"], "url": f"doc:{d['doc_id']}#{c['chunk_id']}",
                    "snippet": c["text"].strip(),
                    "kind": "doc", "doc_name": d["name"]})
    return out


def list_docs() -> list[dict]:
    return [{"doc_id": d["doc_id"], "name": d["name"], "chunks": len(d["chunks"]),
             "chars": d.get("chars", 0), "added": d.get("added")} for d in _load()]


def count() -> int:
    return len(_load())


def chunk_count() -> int:
    return sum(len(d["chunks"]) for d in _load())


def clear() -> int:
    global _docs
    n = len(_load())
    _docs = []
    _save()
    return n
